import io, os
from docx import Document
from docx.shared import Inches
from docxtpl import DocxTemplate, InlineImage
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
from .storage import storage

def make_placeholder_template(path: str, title: str):
    """Placeholder letterhead. Replace with a real .docx whose header holds the
    letterhead; the body only needs the {{ }} tags below, so formatting is untouched."""
    d = Document()
    h = d.sections[0].header.paragraphs[0]; h.text = title; h.alignment = 1
    for t in ["From:\n{{ sender }}", "{{ through_block }}", "To:\n{{ recipient }}",
              "Subject: {{ subject }}", "Respected Sir/Madam,", "{{ body }}",
              "Thank you.", "{%p for s in signatories %}", "{{ s.image }}",
              "{{ s.name }}\n{{ s.position }}, {{ s.organization }}", "{%p endfor %}"]:
        d.add_paragraph(t)
    os.makedirs(os.path.dirname(path), exist_ok=True); d.save(path)

def party(p) -> str:
    return f"{p['name']}\n{p['position']}\n{p['organization']}"

def build_docx(letter, template, sigs) -> bytes:
    tpl = DocxTemplate(storage.path(template.file_path))
    ctx = {
        "sender": party(letter.sender), "recipient": party(letter.recipient),
        "through_block": f"Through:\n{party(letter.through)}" if letter.through else "",
        "subject": letter.subject, "body": letter.body,
        "signatories": [{
            "name": s.name, "position": s.position, "organization": s.organization,
            "image": InlineImage(tpl, storage.path(s.signature_path), width=Inches(1.4)) if s.signature_path else "",
        } for s in sigs],
    }
    tpl.render(ctx)
    buf = io.BytesIO(); tpl.save(buf); return buf.getvalue()

def build_pdf(letter, template, sigs) -> bytes:
    ss = getSampleStyleSheet(); n = ss["BodyText"]
    head = ParagraphStyle("h", parent=ss["Title"], fontSize=16)
    br = lambda s: s.replace("\n", "<br/>")
    el = [Paragraph(template.name, head), Spacer(1, 18),
          Paragraph("<b>From:</b><br/>" + br(party(letter.sender)), n), Spacer(1, 8)]
    if letter.through:
        el += [Paragraph("<b>Through:</b><br/>" + br(party(letter.through)), n), Spacer(1, 8)]
    el += [Paragraph("<b>To:</b><br/>" + br(party(letter.recipient)), n), Spacer(1, 12),
           Paragraph(f"<b>Subject: {letter.subject}</b>", n), Spacer(1, 12),
           Paragraph("Respected Sir/Madam,", n), Spacer(1, 6)]
    for para in letter.body.split("\n\n"):
        el += [Paragraph(br(para), n), Spacer(1, 6)]
    el += [Paragraph("Thank you.", n), Spacer(1, 20)]
    for s in sigs:
        if s.signature_path:
            el.append(Image(storage.path(s.signature_path), width=100, height=40, kind="proportional"))
        el += [Paragraph(f"<b>{s.name}</b><br/>{s.position}, {s.organization}", n), Spacer(1, 12)]
    buf = io.BytesIO(); SimpleDocTemplate(buf, pagesize=A4, topMargin=50).build(el); return buf.getvalue()
